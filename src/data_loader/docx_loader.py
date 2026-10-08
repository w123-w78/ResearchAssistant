# -*- coding: utf-8 -*-
"""
模块一 5.1.1：Word 文档加载器（.docx / .doc）

主路径：python-docx（段落 + 表格 + 标题层级）
兜底路径：直接解析 word/document.xml（零依赖，保证离线环境也能用）
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Tuple
from xml.etree import ElementTree as ET

from .base import BaseLoader, Document, Page, build_metadata, make_doc_id

try:
    from ..utils.logger import get_logger

    logger = get_logger("research_assistant.loader.docx")
except Exception:  # pragma: no cover
    import logging

    logger = logging.getLogger("docx_loader")

_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class DocxLoader(BaseLoader):
    """Word 文档加载器。docx 原生支持；doc 为旧二进制格式，需提示用户先转换。"""

    supported_exts = ["docx"]

    def __init__(self, clean: bool = True, keep_formula: bool = True,
                 include_tables: bool = True, split_by_heading: bool = False):
        super().__init__(clean=clean, keep_formula=keep_formula)
        self.include_tables = include_tables
        # 按一级标题拆分为多个"逻辑页"，便于引用溯源定位到章节
        self.split_by_heading = split_by_heading

    def load(self, path: str) -> Document:
        ext = Path(path).suffix.lower()
        if ext == ".doc":
            raise ValueError(
                "不支持旧版 .doc 二进制格式，请先另存为 .docx"
                "（Word: 文件 -> 另存为 -> Word 文档(*.docx)）"
            )
        try:
            import docx  # python-docx  # noqa: F401

            pages, meta = self._load_with_docx(path)
            parser = "python-docx"
        except ImportError:
            logger.warning("python-docx 未安装，使用内置 XML 解析兜底")
            pages, meta = self._load_with_xml(path)
            parser = "builtin-xml"

        pages = self._postprocess(pages)
        meta.update(build_metadata(path, {"parser": parser}))
        return Document(
            doc_id=make_doc_id(path),
            title=meta.pop("heading_title", "") or Path(path).stem,
            source=str(Path(path).resolve()),
            file_type="docx",
            pages=pages,
            metadata=meta,
        )

    # ---------------- python-docx 路径 ----------------
    def _load_with_docx(self, path: str) -> Tuple[List[Page], Dict[str, Any]]:
        import docx as docx_mod

        d = docx_mod.Document(path)
        meta: Dict[str, Any] = {}
        blocks: List[str] = []
        heading_title = ""

        core = getattr(d, "core_properties", None)
        if core is not None:
            meta["author"] = getattr(core, "author", "") or ""
            meta["created"] = str(getattr(core, "created", "") or "")

        # 段落
        for p in d.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name or "") if p.style is not None else ""
            if style.lower().startswith("heading 1") or style.lower() == "title":
                if not heading_title:
                    heading_title = text[:200]
                blocks.append(f"\n## {text}\n")
            elif style.lower().startswith("heading"):
                blocks.append(f"\n### {text}\n")
            else:
                blocks.append(text)

        # 表格
        if self.include_tables and d.tables:
            meta["table_count"] = len(d.tables)
            for tb in d.tables:
                rows = [[c.text.strip() for c in r.cells] for r in tb.rows]
                blocks.append(self._table_to_md(rows))

        meta["heading_title"] = heading_title
        if self.split_by_heading:
            return self._split_by_heading(blocks), meta
        return [Page(page=1, text="\n".join(blocks), metadata={"source": "python-docx"})], meta

    # ---------------- 内置 XML 兜底 ----------------
    def _load_with_xml(self, path: str) -> Tuple[List[Page], Dict[str, Any]]:
        meta: Dict[str, Any] = {"heading_title": ""}
        blocks: List[str] = []
        with zipfile.ZipFile(path) as z:
            if "word/document.xml" not in z.namelist():
                raise ValueError("不是有效的 .docx 文件（缺少 word/document.xml）")
            xml = z.read("word/document.xml")
        root = ET.fromstring(xml)
        body = root.find(f"{_W_NS}body")
        if body is None:
            return [Page(page=1, text="")], meta

        for para in body.iter(f"{_W_NS}p"):
            text = "".join(t.text or "" for t in para.iter(f"{_W_NS}t")).strip()
            if not text:
                continue
            ppr = para.find(f"{_W_NS}pPr")
            style = ""
            if ppr is not None:
                st = ppr.find(f"{_W_NS}pStyle")
                if st is not None:
                    style = st.get(f"{_W_NS}val", "") or ""
            if style.lower().startswith("heading1") or style.lower() == "title":
                if not meta["heading_title"]:
                    meta["heading_title"] = text[:200]
                blocks.append(f"\n## {text}\n")
            elif style.lower().startswith("heading"):
                blocks.append(f"\n### {text}\n")
            else:
                blocks.append(text)

        # 表格（简化版：每个 tbl 逐行拼接）
        tbl_count = 0
        for tbl in body.iter(f"{_W_NS}tbl"):
            tbl_count += 1
            rows = []
            for tr in tbl.iter(f"{_W_NS}tr"):
                cells = ["".join(t.text or "" for t in tc.iter(f"{_W_NS}t")).strip()
                         for tc in tr.iter(f"{_W_NS}tc")]
                rows.append(cells)
            if rows:
                blocks.append(self._table_to_md(rows))
        meta["table_count"] = tbl_count

        if self.split_by_heading:
            return self._split_by_heading(blocks), meta
        return [Page(page=1, text="\n".join(blocks), metadata={"source": "builtin-xml"})], meta

    # ---------------- 工具 ----------------
    @staticmethod
    def _split_by_heading(blocks: List[str]) -> List[Page]:
        """按 '## ' 一级标题切分为多个逻辑页，便于引用定位到章节。"""
        pages: List[Page] = []
        cur: List[str] = []
        idx = 1
        for b in blocks:
            if b.strip().startswith("## "):
                if cur:
                    pages.append(Page(page=idx, text="\n".join(cur)))
                    idx += 1
                    cur = []
            cur.append(b)
        if cur:
            pages.append(Page(page=idx, text="\n".join(cur)))
        return pages or [Page(page=1, text="")]

    @staticmethod
    def _table_to_md(rows: List[List[str]]) -> str:
        if not rows:
            return ""
        width = max(len(r) for r in rows)
        norm = [r + [""] * (width - len(r)) for r in rows]
        lines = ["| " + " | ".join(norm[0]) + " |", "|" + "|".join(["---"] * width) + "|"]
        for r in norm[1:]:
            lines.append("| " + " | ".join(r) + " |")
        return "\n".join(lines)
