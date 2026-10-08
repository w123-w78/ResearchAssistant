# -*- coding: utf-8 -*-
"""把 docs/项目交付书.md 按 项目交付模板.docx 的排版规则转换为 Word 文档。

用法：
    python scripts/md_to_docx.py [--md docs/项目交付书.md] [--out 项目交付书.docx]
    python scripts/md_to_docx.py --template "C:/Users/34288/Desktop/生产实习/项目交付模板.docx"

排版规则（取自模板实测）：
    页面 A4，上下 25mm、左右 32mm
    封面：仿宋 / 黑体 15pt / 隶书 36pt
    正文：宋体 12pt，首行缩进 2 字符，1.5 倍行距
    标题：章 黑体 16pt 居中；节 黑体 14pt；小节 黑体 12pt 加粗
    表格：Table Grid，表头加粗并加底纹
"""
import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---------- 排版常量（与模板一致） ----------
PAGE_TOP, PAGE_BOTTOM = Cm(2.5), Cm(2.5)
PAGE_LEFT, PAGE_RIGHT = Cm(3.2), Cm(3.2)
BODY_FONT, BODY_SIZE = "宋体", Pt(12)
HEAD_FONT = "黑体"
MONO_FONT, MONO_SIZE = "Consolas", Pt(9.5)
CODE_FONT_CN = "等线"

HEADING_SPEC = {
    1: (HEAD_FONT, Pt(16), True, WD_ALIGN_PARAGRAPH.CENTER),
    2: (HEAD_FONT, Pt(14), True, WD_ALIGN_PARAGRAPH.LEFT),
    3: (HEAD_FONT, Pt(12), True, WD_ALIGN_PARAGRAPH.LEFT),
    4: (HEAD_FONT, Pt(12), True, WD_ALIGN_PARAGRAPH.LEFT),
}


def set_run(run, cn_font=BODY_FONT, en_font="Times New Roman", size=BODY_SIZE,
            bold=False, color=None):
    run.font.name = en_font
    run.font.size = size
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn_font)
    rfonts.set(qn("w:ascii"), en_font)
    rfonts.set(qn("w:hAnsi"), en_font)


def shade(element, hex_fill):
    """给段落或单元格加底纹。"""
    if hasattr(element, "_tc"):          # 表格单元格
        pr = element._tc.get_or_add_tcPr()
    else:                                 # 段落
        pr = element._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def add_page_number_footer(section):
    """页脚插入居中页码（域代码）。"""
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run()
    set_run(r1, BODY_FONT, size=Pt(10))
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r1._element.append(fld)


INLINE_BOLD = re.compile(r"\*\*(.+?)\*\*")
INLINE_CODE = re.compile(r"`([^`]+)`")


def add_inline(p, text, cn_font=BODY_FONT, size=BODY_SIZE, base_bold=False):
    """解析行内 **粗体** 与 `代码`，逐段写入 run。"""
    text = text.replace("\\", "")
    parts = INLINE_BOLD.split(text)
    for i, seg in enumerate(parts):
        is_bold = (i % 2 == 1)
        for j, sub in enumerate(INLINE_CODE.split(seg)):
            if not sub:
                continue
            is_code = (j % 2 == 1)
            r = p.add_run(sub)
            set_run(r, cn_font=MONO_FONT if is_code else cn_font,
                    en_font=MONO_FONT if is_code else "Times New Roman",
                    size=MONO_SIZE if is_code else size,
                    bold=base_bold or is_bold)


def strip_md(text):
    """表格单元格里去掉标记，只留纯文本。"""
    return INLINE_CODE.sub(r"\1", INLINE_BOLD.sub(r"\1", text)).strip()


# ---------- 块级渲染 ----------
def add_heading(doc, text, level):
    font, size, bold, align = HEADING_SPEC.get(level, HEADING_SPEC[3])
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.25
    if level == 1:
        pf.page_break_before = True
    add_inline(p, text, cn_font=font, size=size, base_bold=bold)
    return p


def add_body(doc, text, indent=True, cn_font=BODY_FONT, size=BODY_SIZE,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(3)
    if indent:
        pf.first_line_indent = Pt(size.pt * 2)
    p.alignment = align
    add_inline(p, text, cn_font=cn_font, size=size)
    return p


def join_lines(buf):
    """合并 Markdown 软换行：中文直接拼接，英文/数字之间补一个空格。"""
    out = buf[0]
    for seg in buf[1:]:
        if (out and seg and out[-1].isascii() and out[-1].isalnum()
                and seg[0].isascii() and seg[0].isalnum()):
            out += " " + seg
        else:
            out += seg
    return out


def add_toc_entry(doc, text, level):
    """目录条目。level=0 章（黑体加粗顶格），level=1 节（宋体缩进）。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.4
    pf.space_after = Pt(1)
    pf.left_indent = Pt(0 if level == 0 else 26)
    if level == 0:
        add_inline(p, text, cn_font=HEAD_FONT, size=Pt(12), base_bold=True)
    else:
        add_inline(p, text, cn_font=BODY_FONT, size=Pt(11))
    return p


def add_bullet(doc, text, ordered=False, index=1):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Pt(24)
    pf.line_spacing = 1.5
    pf.space_after = Pt(2)
    prefix = f"{index}. " if ordered else "· "
    r = p.add_run(prefix)
    set_run(r)
    add_inline(p, text)
    return p


def add_quote(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Pt(24)
    pf.line_spacing = 1.5
    pf.space_after = Pt(3)
    add_inline(p, text, size=Pt(11))
    for r in p.runs:
        r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    shade(p, "F5F5F5")
    return p


def add_code(doc, lines):
    for ln in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Pt(18)
        pf.line_spacing = 1.15
        pf.space_after = Pt(0)
        r = p.add_run(ln if ln else " ")
        set_run(r, cn_font=CODE_FONT_CN, en_font=MONO_FONT, size=MONO_SIZE)
        shade(p, "F4F4F4")


def add_image_or_placeholder(doc, alt, rel_path, base_dir):
    img = (base_dir / rel_path) if rel_path else None
    if img and img.exists():
        try:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(img), width=Cm(14))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cap.add_run(alt)
            set_run(r, size=Pt(10.5))
            return
        except Exception as exc:  # 图片损坏等情况降级为占位
            print(f"  [警告] 图片插入失败 {img}: {exc}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"【待插入截图】{alt}　（文件：{rel_path}）")
    set_run(r, cn_font=HEAD_FONT, size=Pt(11),
            color=RGBColor(0x88, 0x55, 0x00))
    shade(p, "FFF6E5")
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)


def add_table(doc, rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci in range(ncol):
            txt = strip_md(row[ci]) if ci < len(row) else ""
            cell = cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.line_spacing = 1.15
            pf.space_after = Pt(2)
            if ri == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(txt)
                set_run(r, cn_font=HEAD_FONT, size=Pt(10.5), bold=True)
                shade(cell, "EFEFEF")
            else:
                r = p.add_run(txt)
                set_run(r, size=Pt(10.5))
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


# ---------- Markdown 解析 ----------
def parse_and_render(doc, md_path):
    base_dir = md_path.parent
    lines = md_path.read_text(encoding="utf-8").split("\n")
    i, n = 0, len(lines)
    first_h1 = True
    toc_mode = False              # 是否处于"目录"区块内
    while i < n:
        line = lines[i]
        s = line.strip()

        # 代码块
        if s.startswith("```"):
            lang = s[3:].strip()
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i].rstrip())
                i += 1
            i += 1
            if lang:
                tag = doc.add_paragraph()
                tag.paragraph_format.space_after = Pt(2)
                r = tag.add_run(f"[{lang}]" if lang != "mermaid" else "[架构图 · Mermaid 源码]")
                set_run(r, cn_font=HEAD_FONT, size=Pt(10),
                        color=RGBColor(0x66, 0x66, 0x66))
            add_code(doc, buf)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        if not s:
            i += 1
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}|\*{3,}", s):
            toc_mode = False
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            level = len(m.group(1))
            title = m.group(2).strip()
            # 只有紧跟 "## 目录" 的列表才按目录层级渲染，其余标题一律退出目录模式
            toc_mode = (level == 2 and title == "目录")
            p = add_heading(doc, title, min(level, 4))
            if level == 1 and first_h1:
                p.paragraph_format.page_break_before = False
                first_h1 = False
            i += 1
            continue

        # 图片
        m = re.match(r"^!\[([^\]]*)\]\(([^)]*)\)$", s)
        if m:
            add_image_or_placeholder(doc, m.group(1), m.group(2), base_dir)
            i += 1
            continue

        # 表格
        if s.startswith("|") and i + 1 < n and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            rows = []
            header = [c.strip() for c in s.strip("|").split("|")]
            rows.append(header)
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, rows)
            continue

        # 引用
        if s.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            add_quote(doc, " ".join(x for x in buf if x))
            continue

        # 有序列表
        m = re.match(r"^(\d+)[.)]\s+(.*)$", s)
        if m:
            add_bullet(doc, m.group(2), ordered=True, index=int(m.group(1)))
            i += 1
            continue

        # 无序列表
        if s.startswith("- ") or s.startswith("* "):
            if toc_mode:
                indent = len(line) - len(line.lstrip())   # 用原始行算缩进
                add_toc_entry(doc, s[2:], 1 if indent >= 2 else 0)
            else:
                add_bullet(doc, s[2:])
            i += 1
            continue

        # 普通正文：合并连续软换行为一个段落（Markdown 语义）
        buf = [s]
        i += 1
        while i < n:
            s2 = lines[i].strip()
            if not s2:
                i += 1
                break
            if (s2[0] in "#|>-*" or s2.startswith("```") or s2.startswith("![")
                    or re.fullmatch(r"-{3,}|\*{3,}", s2)
                    or re.match(r"^\d+[.)]\s+", s2)
                    or re.match(r"^\[\d+\]", s2)):   # 参考文献条目不参与合并
                break
            buf.append(s2)
            i += 1
        text = join_lines(buf)
        if re.match(r"^\[\d+\]", text):          # 参考文献条目：悬挂缩进
            p = add_body(doc, text, indent=False)
            pf = p.paragraph_format
            pf.left_indent = Pt(22)
            pf.first_line_indent = Pt(-22)
        else:
            add_body(doc, text)


# ---------- 封面（沿用模板格式） ----------
def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("南京农业大学课程考核情况表")
    set_run(r, cn_font="仿宋", size=Pt(18), bold=True)
    p.paragraph_format.space_after = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("2026 — 2027 学年第 1 学期")
    set_run(r, cn_font=HEAD_FONT, size=Pt(15))
    p.paragraph_format.space_after = Pt(14)

    for text in [
        "学    院：　智慧农业学院（人工智能学院）",
        "课程名称：　生产实习　　　　学分：　1",
        "考核方式：　课程报告",
        "考核内容：　智能科研助理 —— 基于 RAG + Agent 的论文知识库问答系统",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.8
        r = p.add_run(text)
        set_run(r, cn_font=HEAD_FONT, size=Pt(15))

    doc.add_paragraph()
    for text in ["成绩：", "教师签名：", "时间：　2026 年 12 月 30 日"]:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.8
        r = p.add_run(text)
        set_run(r, cn_font=HEAD_FONT, size=Pt(15))

    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("生产实习课程报告")
    set_run(r, cn_font="隶书", size=Pt(36), bold=True)
    p.paragraph_format.space_after = Pt(28)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    r = p.add_run("题　　目：　智能科研助理 —— 基于 RAG + Agent 的论文知识库问答系统")
    set_run(r, cn_font=BODY_FONT, size=Pt(16), bold=True)

    for k in range(1, 5):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 2.0
        r = p.add_run(f"成 员 {k}：　班级 ____________　学号 ____________　姓名 ____________")
        set_run(r, cn_font=BODY_FONT, size=Pt(16), bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--md", default="docs/项目交付书.md")
    ap.add_argument("--out", default=None)
    ap.add_argument("--template", default="C:/Users/34288/Desktop/生产实习/项目交付模板.docx")
    args = ap.parse_args()

    root = Path(__file__).resolve().parent.parent
    md_path = Path(args.md)
    if not md_path.is_absolute():
        md_path = root / md_path
    out_path = Path(args.out) if args.out else (Path(args.template).parent / "项目交付书.docx")

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = BODY_SIZE
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = PAGE_TOP, PAGE_BOTTOM
    sec.left_margin, sec.right_margin = PAGE_LEFT, PAGE_RIGHT
    add_page_number_footer(sec)

    add_cover(doc)
    parse_and_render(doc, md_path)

    doc.save(str(out_path))
    print(f"已生成：{out_path}")
    print(f"段落数 {len(doc.paragraphs)}　表格数 {len(doc.tables)}")


if __name__ == "__main__":
    sys.exit(main())
